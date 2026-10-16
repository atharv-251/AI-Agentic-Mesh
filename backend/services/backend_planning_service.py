import io
from datetime import datetime, timezone

from docx import Document
from docx.shared import Inches, Pt
from sqlalchemy.orm import Session

from backend.ai.agents import KnowledgeGraphAgent, KTLevelAgent, ProjectProfileAgent
from backend.models import BackendPlan, KTLevelEvaluation, KTSession, KnowledgeNode, ProjectProfile, RawExtraction, Stakeholder, Transition
from backend.schemas.backend_plan import PlanningProposal
from backend.services.capacity_service import CapacityService
from backend.services.scheduling_service import SchedulingService
from backend.services.validation_service import ValidationService
from backend.services.tracking_service import sync_session_activities


def execute_plan(db: Session, plan: BackendPlan, proposal: PlanningProposal) -> dict:
    transition = db.get(Transition, plan.transition_id)
    if transition is None:
        raise ValueError("Transition no longer exists")
    if any(db.query(model).filter_by(transition_id=transition.id).first() is not None
           for model in (ProjectProfile, KnowledgeNode, Stakeholder, KTSession, RawExtraction)):
        raise ValueError("This KT already has operational inputs. Use a new workspace; existing records are never overwritten")
    if not transition.start_date or not transition.end_date or transition.end_date < transition.start_date:
        raise ValueError("Set a valid transition date range before execution")
    if not {"sme", "receiver"} <= {person.role for person in proposal.people}:
        raise ValueError("Provide at least one reviewed SME and receiver with explicit KT levels")
    if any(topic.category == "unclassified" for topic in proposal.topics):
        raise ValueError("Classify all Amplify topics before operational calculation")
    if any(abs(topic.hours * 10 - round(topic.hours * 10)) > 1e-6 for topic in proposal.topics):
        raise ValueError("Topic hours must use tenths of an hour, matching the existing hierarchy engine")
    cap = CapacityService.calculate_transition_capacity_targets(
        transition.total_duration_days, transition.shadow_days,
        transition.reverse_shadow_days, transition.daily_kt_hours,
    )
    if cap["target_capacity_hours"] <= 0:
        raise ValueError("Transition settings must provide positive KT capacity")
    transition.available_kt_days = cap["available_kt_days"]
    transition.target_capacity_hours = cap["target_capacity_hours"]
    raw = {
        "project_name": proposal.project_name, "project_category": proposal.category,
        "source_files": [source["filename"] for source in plan.source_files],
        "technology_stack": proposal.technology_stack, "environments": proposal.environments,
        "support_model": proposal.support_model, "integrations": proposal.integrations,
        "dependencies": proposal.dependencies, "risks": proposal.risks,
        "evidence_gaps": proposal.uncertainties,
        "applications": [{
            "application_name": proposal.project_name,
            "business_purpose": proposal.business_purpose, "criticality": "Not supplied",
            "topics": [{
                "topic": topic.name, "category": topic.category,
                "subtopics": ["Reviewed scope"], "evidence": [topic.evidence],
            } for topic in proposal.topics],
        }],
        "kt_duration_analysis": {"topic_effort_breakdown": [
            {"topic": topic.name, "estimated_hours": topic.hours} for topic in proposal.topics
        ]},
    }
    db.add(RawExtraction(transition_id=transition.id, raw_json_payload=raw, normalized_payload=raw))
    db.commit()
    profile = ProjectProfileAgent.run(db, transition.id)
    profile.intended_levels = proposal.intended_levels
    db.commit()
    for person in proposal.people:
        db.add(Stakeholder(
            transition_id=transition.id, name=person.name, role=person.role,
            level=person.level, email=person.email,
        ))
    db.commit()
    KnowledgeGraphAgent.run(db, transition.id)
    KTLevelAgent.run(db, transition.id)
    # Do not inflate a gap using the legacy generic expansion topics.
    capacity = CapacityService.evaluate_capacity_balance(db, transition.id)
    sessions = SchedulingService.auto_schedule_sessions(
        db, transition.id, daily_max_hours=transition.daily_kt_hours,
    )
    sync_session_activities(transition.id, [], db)
    db.commit()
    validation = ValidationService.run_full_validation(db, transition.id).model_dump(mode="json")
    hierarchy = db.query(KnowledgeNode).filter_by(transition_id=transition.id).order_by(KnowledgeNode.order_index).all()
    evaluations = db.query(KTLevelEvaluation).filter_by(transition_id=transition.id).all()
    scheduled_hours = round(sum(session.duration_hours for session in sessions), 2)
    blockers = list(proposal.uncertainties)
    if capacity["gap_hours"] > 0:
        blockers.append("Reviewed topics do not cover target capacity; add evidence-backed scope instead of inventing topics")
    if abs(scheduled_hours - capacity["generated_hours"]) > 0.01:
        blockers.append("Scheduled hours do not cover all reviewed topic effort")
    if not sessions:
        blockers.append("No schedulable sessions were generated")
    for check in validation["checks"]:
        if not check["passed"]:
            blockers.append(check["message"])
    blockers.append("Calendar and leave completeness must be reviewed; no external calendar synchronization is implied")
    transition.status = "backend_draft"
    result = {
        "execution": "original_python_services", "release_status": "draft_not_released",
        "capacity": capacity, "validation": validation, "blockers": blockers,
        "scheduled_hours": scheduled_hours,
        "sessions": [{
            "id": session.id, "title": session.session_title, "level": session.level,
            "date": session.scheduled_date.isoformat(), "start": session.start_time.isoformat(),
            "end": session.end_time.isoformat(), "hours": session.duration_hours,
            "sme_id": session.sme_id, "receiver_id": session.receiver_id,
            "sme_name": db.get(Stakeholder, session.sme_id).name if session.sme_id else "Unassigned",
            "receiver_name": db.get(Stakeholder, session.receiver_id).name if session.receiver_id else "Unassigned",
            "conflicts": session.conflict_flags,
        } for session in sessions],
        "people": [person.model_dump() for person in proposal.people],
        "topics": [topic.model_dump() for topic in proposal.topics],
        "hierarchy": [{
            "id": node.id, "parent_id": node.parent_id, "type": node.node_type,
            "name": node.name, "hours": node.estimated_hours, "evidence": node.evidence_references,
        } for node in hierarchy],
        "level_evaluations": [{
            "node_id": item.node_id, "scope": item.level_scope,
            "objective": item.learning_objective, "outcome": item.expected_outcome,
        } for item in evaluations],
        "settings": {
            "start_date": transition.start_date.isoformat(), "end_date": transition.end_date.isoformat(),
            "total_days": transition.total_duration_days, "shadow_days": transition.shadow_days,
            "reverse_shadow_days": transition.reverse_shadow_days, "daily_hours": transition.daily_kt_hours,
            "sme_country": transition.sme_country, "receiver_country": transition.receiver_country,
        },
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "messages_sent": False, "tracking_status": "awaiting_actual_transcripts",
    }
    return result


def create_report(plan: BackendPlan, proposal: PlanningProposal, result: dict) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.header.paragraphs[0].text = "KT PLANNER | CONTROLLED DRAFT"
    section.footer.paragraphs[0].text = f"Plan {plan.id} | Not a release authorization"
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    section.top_margin = Inches(0.7)
    doc.add_heading(f"Knowledge Transfer Plan\n{proposal.project_name}", 0)
    doc.add_paragraph("DRAFT - HUMAN REVIEW REQUIRED. No invitations or release approval are implied.")
    doc.add_paragraph(f"Prepared: {result['generated_at']}\nReviewed inputs by: {plan.reviewed_by}")
    doc.add_heading("1. Executive summary and application profile", 1)
    doc.add_paragraph(proposal.business_purpose or "Business purpose not supplied.")
    for key in ("category", "intended_levels", "technology_stack", "environments", "support_model", "integrations", "dependencies", "risks"):
        doc.add_paragraph(f"{key.replace('_', ' ').title()}: {getattr(proposal, key)}")
    doc.add_heading("2. Transition settings", 1)
    for key, value in result["settings"].items():
        doc.add_paragraph(f"{key.replace('_', ' ').title()}: {value}")
    doc.add_heading("3. SMEs and receivers", 1)
    for person in proposal.people:
        doc.add_paragraph(f"{person.name} | {person.role} | {person.level} | {person.email or 'Email not supplied'}")
        doc.add_paragraph(f"Evidence: {person.evidence}")
    doc.add_heading("4. Knowledge scope and effort", 1)
    for topic in proposal.topics:
        doc.add_paragraph(f"{topic.name} | {topic.category} | {topic.hours:g} hours")
        doc.add_paragraph(f"Evidence: {topic.evidence}")
    doc.add_heading("Operational hierarchy and learning outcomes", 2)
    for node in result["hierarchy"]:
        doc.add_paragraph(f"{node['type'].title()}: {node['name']} | {node['hours']:g}h | Record {node['id']}")
    node_names = {node["id"]: node["name"] for node in result["hierarchy"]}
    for evaluation in result["level_evaluations"]:
        doc.add_paragraph(f"{node_names[evaluation['node_id']]} | {evaluation['scope']}")
        doc.add_paragraph(f"Objective: {evaluation['objective']}\nExpected outcome: {evaluation['outcome']}")
    doc.add_heading("5. Deterministic capacity calculation", 1)
    for key, value in result["capacity"].items():
        doc.add_paragraph(f"{key.replace('_', ' ').title()}: {value}")
    doc.add_heading("6. Proposed session schedule", 1)
    for session in result["sessions"]:
        doc.add_paragraph(
            f"{session['date']} {session['start']}-{session['end']} | {session['title']} | "
            f"{session['level']} | {session['hours']:g}h\n"
            f"SME: {session['sme_name']} | Receiver: {session['receiver_name']}\n"
            f"Conflicts: {session['conflicts'] or 'None detected from supplied records'}"
        )
    doc.add_heading("7. Validation results", 1)
    for check in result["validation"]["checks"]:
        doc.add_paragraph(f"{'PASS' if check['passed'] else 'REVIEW'} | {check['check_name']}: {check['message']}")
    doc.add_heading("8. Blockers, assumptions and review actions", 1)
    for blocker in result["blockers"]:
        doc.add_paragraph(blocker, style="List Bullet")
    doc.add_heading("9. Tracking and handover", 1)
    doc.add_paragraph("Awaiting actual KT sessions, transcripts, delivered documents and acceptance evidence. No completion or readiness has been inferred.")
    doc.add_heading("10. Provenance and execution boundary", 1)
    provenance = result.get("amplify_provenance")
    if provenance:
        doc.add_paragraph(f"Mandatory Amplify workflow 167 source run: {provenance['run_id']}. Source transfer: manually imported, unverified snapshot. Human-reviewed normalization was used by Python services for deterministic operational calculations; the original Amplify outputs remain separate and unchanged.")
    else:
        doc.add_paragraph("Historical backend-only execution. This is not an Amplify-generated plan.")
    doc.add_paragraph("Python services saved the operational records and calculated validation results. Delivery of this draft does not approve the transition or send calendar invitations.")
    doc.add_paragraph(f"KT: {plan.transition_id}\nExecution: {plan.id}")
    for source in plan.source_files:
        doc.add_paragraph(f"{source['filename']} | SHA-256: {source['sha256']} | {source['characters']} extracted characters")
    if result.get("amplify_agent_outputs"):
        doc.add_heading("Appendix: original seven-agent Amplify outputs", 1)
        doc.add_paragraph("Preserved exactly as imported. These source proposals are not replaced by the operational calculation above; differences require human review. Snapshot authenticity is not independently verified.")
        for output in result["amplify_agent_outputs"]:
            doc.add_heading(f"Agent {output['agent_id']} - {output['agent_name'] or 'Workflow agent'}", 2)
            doc.add_paragraph(output["result"])
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()
